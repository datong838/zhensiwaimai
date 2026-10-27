#!/usr/bin/env python3
"""台账写入器 + 轮转排期器（纯本地逻辑，不碰网络、不碰浏览器）。

为什么单独抽出来
----------------
截图/发送依赖 9230 登录态，会反复被外部阻塞；但「哪天该发哪个指标」「哪个店哪个指标
已发过」「重跑会不会写出重复行」这些是**确定性逻辑**，必须能离线自证。
把它们和采集器解耦，登录一恢复就只剩「截一张图」这一步。

两个子命令
----------
  record  写入一条反馈记录（幂等），同步宽表打勾 + 重算已完成数
  plan    算指定日期该发哪个指标、还差哪些店、哪些店要顺延

幂等语义（对齐两处权威口径）
----------------------------
* 幂等键 = (店铺id, 指标序号, 记录所属月份)
* 已有 `verified` / `unknown` ⇒ 视为**已发过**，重跑返回 `already`，绝不重复外发
  （内存 #79「一事一发」：#75 `unknown` 也不重发；重复外发是事故，漏发只是延迟）
* 已有 `blocked` / `no-data` ⇒ 允许补发（当日/次日顺延），不算重复

用法
----
  python scripts/ledger_ops.py plan --date 2026-10-21
  python scripts/ledger_ops.py record --store-id 5296218123 --metric M01 \\
      --result verified --group "XX店福利群" --shot /path/a.png --apply
  默认 dry-run，只有带 --apply 才真正落盘。
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys
from datetime import date, datetime

from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent
WF = HERE.parent
CATALOG = WF / "docs" / "metrics_catalog.json"

RECORD_SHEET = "当月反馈记录"
PROGRESS_SHEET = "指标进度"
POOL_SHEET = "店铺池"

RECORD_COLS = ["日期", "轮次", "店铺id", "店铺名称", "指标序号", "指标名称",
               "企业微信群", "消息结果", "回执／消息ID", "截图路径", "备注"]
PROGRESS_HEAD = ["店铺id", "店铺名称", "企业微信群名称"]

# 已发过（重跑不得再发）
DONE_RESULTS = {"verified", "skipped-already"}
# 已发出但无回执：同样按已发处理，绝不重发（#75 / #79）
SENT_UNKNOWN = {"unknown"}
# 可重试
RETRYABLE = {"blocked", "no-data"}


# --------------------------------------------------------------------------- 基础
def default_ledger() -> pathlib.Path:
    allx = list((WF / "docs").glob("全店铺运营指标商家反馈台账_*.xlsx"))
    cands = [p for p in allx if "模板" not in p.name and re.search(r"_\d{4}-\d{2}\.xlsx$", p.name)]
    if not cands:
        raise SystemExit(f"未找到当月实盘台账；候选={[p.name for p in allx]}")
    return max(cands, key=lambda p: p.stat().st_mtime)


def load_catalog(path: pathlib.Path = CATALOG) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def registered_metrics(cat: dict) -> list[dict]:
    return sorted(cat.get("metrics", []), key=lambda m: m["id"])


def month_of(d: date) -> str:
    return d.strftime("%Y-%m")


def parse_date(s: str) -> date:
    s = (s or "").strip()
    for f in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], f).date()
        except ValueError:
            continue
    raise SystemExit(f"无法解析日期：{s!r}（要 YYYY-MM-DD）")


def norm_id(v) -> str:
    """店铺id 全库统一成字符串，避免 123 / '123' / 123.0 三种形态破坏幂等。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ------------------------------------------------------------------- 幂等检查
def existing_keys(ws, month: str) -> dict[tuple[str, str], str]:
    """返回 {(店铺id, 指标序号): 最强结果}，只统计同月记录。"""
    head = [c.value for c in ws[1]]
    try:
        i_date = head.index("日期")
        i_sid = head.index("店铺id")
        i_mid = head.index("指标序号")
        i_res = head.index("消息结果")
    except ValueError as e:
        raise SystemExit(f"『{RECORD_SHEET}』表头不符合预期：{head}") from e

    best: dict[tuple[str, str], str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[i_sid] is None:
            continue
        try:
            d = row[i_date]
            if isinstance(d, datetime):
                d = d.date()
            elif isinstance(d, str):
                d = parse_date(d)
            elif not isinstance(d, date):
                continue
        except SystemExit:
            continue
        if month_of(d) != month:
            continue
        k = (norm_id(row[i_sid]), str(row[i_mid] or "").strip())
        res = str(row[i_res] or "").strip()
        # 结果强度：verified/unknown/skipped-already 都算「已发过」，blocked 最弱
        if st := best.get(k):
            if st in DONE_RESULTS and res in SENT_UNKNOWN:
                best[k] = res
            elif st in DONE_RESULTS | SENT_UNKNOWN:
                continue
            elif res in DONE_RESULTS | SENT_UNKNOWN:
                best[k] = res
            else:
                best[k] = st or res
        else:
            best[k] = res
    return best


# ------------------------------------------------------------------- record 子命令
def cmd_record(args) -> int:
    ledger = args.ledger or default_ledger()
    cat = load_catalog()
    metrics = {m["id"]: m for m in registered_metrics(cat)}

    if args.metric not in metrics:
        raise SystemExit(f"{args.metric} 未在指标字典登记（现有：{sorted(metrics)}），拒绝写入")

    d = parse_date(args.date) if args.date else date.today()
    month = month_of(d)

    wb = load_workbook(ledger)
    ws_rec = wb[RECORD_SHEET]
    ws_prg = wb[PROGRESS_SHEET]

    head = [c.value for c in ws_rec[1]]
    if head[:len(RECORD_COLS)] != RECORD_COLS:
        raise SystemExit(f"『{RECORD_SHEET}』表头与约定不符：{head[:len(RECORD_COLS)]}")

    sid = norm_id(args.store_id)
    key = (sid, args.metric)
    prev = existing_keys(ws_rec, month).get(key)

    if prev in DONE_RESULTS:
        print(json.dumps({"decision": "already", "reason": f"同月已有 {prev} 记录",
                          "store_id": sid, "metric": args.metric, "month": month},
                         ensure_ascii=False))
        return 0
    if prev in SENT_UNKNOWN and not args.force:
        print(json.dumps({"decision": "already", "reason": "同月已有 unknown（已发出，无回执）⇒ 绝不重发",
                          "store_id": sid, "metric": args.metric, "month": month},
                         ensure_ascii=False))
        return 0

    # 店铺名/群名：优先用入参，否则从店铺池兜底
    store_name, group = args.store_name or "", args.group or ""
    if not store_name or not group:
        ws_pool = wb[POOL_SHEET]
        ph = [c.value for c in ws_pool[1]]
        i_id, i_name = ph.index("店铺id"), ph.index("店铺名称")
        i_group = ph.index("企业微信群名称")
        for row in ws_pool.iter_rows(min_row=2, values_only=True):
            if norm_id(row[i_id]) == sid:
                store_name = store_name or str(row[i_name] or "")
                group = group or str(row[i_group] or "")
                break

    if not group and args.require_group:
        print(json.dumps({"decision": "blocked", "reason": "企业微信群名称缺失（fail-closed，不外发）",
                          "store_id": sid, "metric": args.metric}, ensure_ascii=False))
        return 3

    row_idx = ws_rec.max_row + 1
    values = [
        datetime.combine(d, datetime.min.time()).date().isoformat(),
        args.round or "",
        sid,
        store_name,
        args.metric,
        metrics[args.metric]["name"],
        group,
        args.result,
        args.receipt or "",
        args.shot or "",
        args.note or "",
    ]
    for col, v in enumerate(values, start=1):
        ws_rec.cell(row=row_idx, column=col, value=v)

    # 宽表打勾：只有成功类结果才写日期（blocked 不写，否则会假装已完成）
    marked = False
    if args.result in DONE_RESULTS | SENT_UNKNOWN:
        ph = [c.value for c in ws_prg[1]]
        if args.metric not in ph:
            raise SystemExit(f"『{PROGRESS_SHEET}』缺少列 {args.metric}")
        mcol = ph.index(args.metric) + 1
        i_id = ph.index("店铺id")
        target = None
        for r in range(2, ws_prg.max_row + 1):
            if norm_id(ws_prg.cell(row=r, column=i_id + 1).value) == sid:
                target = r
                break
        if target is None:
            print(json.dumps({"decision": "blocked",
                              "reason": "店铺不在『指标进度』表（需先跑 store_pool_to_ledger 装载）",
                              "store_id": sid}, ensure_ascii=False))
            return 4
        if not ws_prg.cell(row=target, column=mcol).value:
            ws_prg.cell(row=target, column=mcol, value=d.isoformat())
            marked = True

    if args.apply:
        wb.save(ledger)
    print(json.dumps({
        "decision": "applied" if args.apply else "dry-run",
        "store_id": sid, "metric": args.metric, "month": month, "date": d.isoformat(),
        "result": args.result, "record_row": row_idx, "wide_marked": marked,
        "ledger": str(ledger),
    }, ensure_ascii=False))
    return 0


# --------------------------------------------------------------------- plan 子命令
def cmd_plan(args) -> int:
    ledger = args.ledger or default_ledger()
    cat = load_catalog()
    reg = registered_metrics(cat)
    if not reg:
        raise SystemExit("指标字典为空，无法排期")

    d = parse_date(args.date) if args.date else date.today()
    month = month_of(d)
    total = int(cat.get("total_metrics") or 20)

    # 固定顺序轮转：本月第 n 天 → 第 (n-1) mod total 个指标
    slot = (d.day - 1) % total
    order = [m["id"] for m in reg]
    if slot >= len(order):
        print(json.dumps({
            "decision": "blocked-not-registered",
            "reason": f"今天轮到第 {slot + 1} 个指标，但只登记了 {len(order)}/{total} 项；"
                      f"请大同先补发规则 {len(order) + 1}",
            "date": d.isoformat(), "slot": slot + 1,
        }, ensure_ascii=False))
        return 3

    today_metric = order[slot]
    wb = load_workbook(ledger)
    ws_prg = wb[PROGRESS_SHEET]
    ws_rec = wb[RECORD_SHEET]
    ph = [c.value for c in ws_prg[1]]
    i_id, i_name = ph.index("店铺id"), ph.index("店铺名称")
    i_group = ph.index("企业微信群名称")
    mcol = ph.index(today_metric) + 1

    todo, done, no_group, retry = [], 0, [], []
    blocked = existing_keys(ws_rec, month)
    for r in range(2, ws_prg.max_row + 1):
        sid = norm_id(ws_prg.cell(row=r, column=i_id + 1).value)
        if not sid:
            continue
        name = str(ws_prg.cell(row=r, column=i_name + 1).value or "")
        group = str(ws_prg.cell(row=r, column=i_group + 1).value or "")
        if ws_prg.cell(row=r, column=mcol).value:
            done += 1
            continue
        item = {"store_id": sid, "store_name": name, "group": group}
        if (sid, today_metric) in blocked:
            retry.append(item)
            continue
        if not group:
            no_group.append(item)
            continue
        todo.append(item)

    out = {
        "date": d.isoformat(),
        "month": month,
        "slot": slot + 1,
        "metric": today_metric,
        "metric_name": next(m["name"] for m in reg if m["id"] == today_metric),
        "registered": f"{len(order)}/{total}",
        "progress": {"done": done, "todo": len(todo), "no_group": len(no_group),
                     "retry_from_record": len(retry)},
        "todo": todo if args.show else todo[:5],
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))
    if no_group:
        print(f"\n⚠️ {len(no_group)} 家店缺『企业微信群名称』⇒ fail-closed 不外发（先补映射）",
              file=sys.stderr)
    return 0


# ----------------------------------------------------------------------------- CLI
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="台账写入器 + 20 项轮转排期器")
    ap.add_argument("--ledger", type=pathlib.Path, default=None)
    sub = ap.add_subparsers(dest="cmd", required=True)

    r = sub.add_parser("record", help="写入一条反馈记录（幂等）")
    r.add_argument("--store-id", required=True)
    r.add_argument("--store-name", default="")
    r.add_argument("--metric", required=True, help="M01..M20")
    r.add_argument("--result", required=True,
                   choices=sorted(DONE_RESULTS | SENT_UNKNOWN | RETRYABLE))
    r.add_argument("--group", default="")
    r.add_argument("--shot", default="")
    r.add_argument("--receipt", default="")
    r.add_argument("--note", default="")
    r.add_argument("--date", default="")
    r.add_argument("--round", default="")
    r.add_argument("--apply", action="store_true", help="真正落盘（默认 dry-run）")
    r.add_argument("--force", action="store_true", help="对 unknown 也强制再写（危险，慎用）")
    r.add_argument("--require-group", action="store_true",
                   help="缺企业微信群名时 fail-closed 拒写")

    p = sub.add_parser("plan", help="算当日指标 + 待发店铺")
    p.add_argument("--date", default="")
    p.add_argument("--show", action="store_true", help="打印全部待发店铺（默认只列前 5）")

    args = ap.parse_args(argv)
    return cmd_record(args) if args.cmd == "record" else cmd_plan(args)


if __name__ == "__main__":
    raise SystemExit(main())
