#!/usr/bin/env python3
"""离线单测（不联网、不碰 9231 岗位、不触发任何外发）。

跑法：
  /Users/ddt/work/projects/ai_agent/.venvs/harmonyos3-phone-control/bin/python \
    -m unittest discover -s tests -v      （在本工作流目录下）
"""
from __future__ import annotations

import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import stats_card  # noqa: E402
import task_workflow as W  # noqa: E402

EFFECT_HEADER = ["门店ID", "门店名称", "合同状态", "经营任务完成率"]


def write_effect(directory: Path, rows: list[list], header: list[str] = EFFECT_HEADER,
                 name: str = "effect.csv") -> Path:
    path = directory / name
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
    return path


def write_alloc(directory: Path, rows: list[tuple[str, str, str]], name: str = "alloc.xlsx") -> Path:
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "门店绑定信息"
    sheet.append(["门店名称", "门店ID", "绑定时间", "运营人"])
    for shop, shop_id, owner in rows:
        sheet.append([shop, shop_id, "2026-01-01T00:00:00", owner])
    path = directory / name
    workbook.save(path)
    return path


class RateScale(unittest.TestCase):
    def test_fraction_column_becomes_percent(self):
        self.assertEqual(W._rate_scale([0.8333, 0.9666]), "fraction")

    def test_percent_column_stays_percent(self):
        self.assertEqual(W._rate_scale([83.33, 96.66]), "percent")

    def test_mixed_column_uses_percent(self):
        """整列只要有一个 >1，就按百分数处理（避免小数/百分数混算）。"""
        self.assertEqual(W._rate_scale([0.83, 91.2]), "percent")

    def test_empty_column_has_no_scale_crash(self):
        self.assertEqual(W._rate_scale([]), "percent")

    def test_parse_rate_handles_text_and_junk(self):
        self.assertAlmostEqual(W._parse_rate("92.5%"), 92.5)
        self.assertAlmostEqual(W._parse_rate(" 88 "), 88.0)
        self.assertIsNone(W._parse_rate(""))
        self.assertIsNone(W._parse_rate("—"))
        self.assertIsNone(W._parse_rate(None))


class MissingColumns(unittest.TestCase):
    """列缺失必须阻断，绝不能产出一张没有完成率的表。"""

    def test_effect_without_rate_column_blocked(self):
        with tempfile.TemporaryDirectory() as workspace:
            write_effect(Path(workspace), [["1", "店", "生效中"]], ["门店ID", "门店名称", "合同状态"])
            with self.assertRaises(ValueError) as ctx:
                W.read_effect(Path(workspace) / "effect.csv")
            self.assertEqual(str(ctx.exception), "effect-missing-task-completion-column")

    def test_effect_without_shop_id_blocked(self):
        with tempfile.TemporaryDirectory() as workspace:
            write_effect(Path(workspace), [["店", "生效中", 0.9]], ["门店名称", "合同状态", "经营任务完成率"])
            with self.assertRaises(ValueError) as ctx:
                W.read_effect(Path(workspace) / "effect.csv")
            self.assertEqual(str(ctx.exception), "effect-missing-shop-id-column")

    def test_alloc_without_owner_blocked(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as workspace:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["门店名称", "门店ID"])
            sheet.append(["店", "1"])
            workbook.save(Path(workspace) / "alloc.xlsx")
            with self.assertRaises(ValueError) as ctx:
                W.read_alloc(Path(workspace) / "alloc.xlsx")
            self.assertEqual(str(ctx.exception), "alloc-missing-required-columns")


class ReportBuilding(unittest.TestCase):
    def downloads(self, workspace: str) -> Path:
        """build_report 只认 <run>/downloads/，夹具必须落在这个子目录里。"""
        directory = Path(workspace) / "downloads"
        directory.mkdir(parents=True, exist_ok=True)
        return directory

    def build(self, workspace: str, day: str = "2026-10-13"):
        return W.build_report(Path(workspace), day)

    def test_vlookup_inserts_owner_after_shop_name(self):
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲店", "生效中", "0.9"]])
            write_alloc(target, [("甲店", "1001", "盯个盯徐从众")])
            report = self.build(workspace)
            self.assertEqual(report["rows"], 1)
            self.assertEqual(report["matched"], 1)
            self.assertEqual(report["unmatched"], 0)
            import openpyxl
            sheet = openpyxl.load_workbook(report["path"]).active
            header = [cell.value for cell in sheet[1]]
            self.assertEqual(header[header.index("运营人")], "运营人")
            self.assertEqual(header.index("运营人"), header.index("门店名称") + 1)
            self.assertEqual(report["name"], "2026-10-13运营-店铺经营任务完成率明细.xlsx")

    def test_unmatched_shop_marked_and_reported(self):
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲店", "生效中", "0.9"],
                                           ["2002", "乙店", "生效中", "0.8"]])
            write_alloc(target, [("甲店", "1001", "刘欢")])
            report = self.build(workspace)
            self.assertEqual((report["rows"], report["matched"], report["unmatched"]), (2, 1, 1))
            self.assertEqual(report["stats"][0]["name"], "刘欢")
            self.assertEqual(report["unmatchedShops"], ["乙店"])

    def test_fraction_rates_are_scaled_in_stats(self):
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲", "生效中", "0.9"],
                                           ["1002", "乙", "生效中", "0.8"]])
            write_alloc(target, [("甲", "1001", "刘欢"), ("乙", "1002", "刘欢")])
            report = self.build(workspace)
            self.assertEqual(report["rateScale"], "fraction")
            self.assertAlmostEqual(report["stats"][0]["avg"], 85.0)

    def test_blank_rate_excluded_from_average_but_group_kept(self):
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲", "生效中", "0.9"],
                                           ["1002", "乙", "生效中", ""]])
            write_alloc(target, [("甲", "1001", "刘欢"), ("乙", "1002", "刘欢")])
            report = self.build(workspace)
            group = report["stats"][0]
            self.assertAlmostEqual(group["avg"], 90.0)
            self.assertEqual((group["shops"], group["counted"], group["missing"]), (2, 1, 1))

    def test_group_without_any_value_is_dropped_from_card(self):
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲", "生效中", ""],
                                           ["1002", "乙", "生效中", "0.9"]])
            write_alloc(target, [("甲", "1001", "空组"), ("乙", "1002", "刘欢")])
            report = self.build(workspace)
            self.assertEqual([item["name"] for item in report["stats"]], ["刘欢"])

    def test_coverage_reports_alloc_only_owners(self):
        """雷洁老师式盲区：有分配、无效果数据 ⇒ 必须显式上报，不能静默消失。"""
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲", "生效中", "0.9"]])
            write_alloc(target, [("甲", "1001", "刘欢"),
                                 ("乙", "1002", "雷洁老师"),
                                 ("丙", "1003", "雷洁老师")])
            report = self.build(workspace)
            coverage = report["coverage"]
            self.assertEqual(coverage["allocOnlyShops"], 2)
            entry = coverage["allocOnlyGroups"][0]
            self.assertEqual((entry["raw"], entry["shops"]), ("雷洁老师", 2))
            self.assertEqual([item["name"] for item in report["stats"]], ["刘欢"])
            # 图尾提示行必须同步生成（大同 10-13 口径）
            self.assertEqual(report["notes"], ["雷洁老师 2 家无数据未计入"])

    def test_zip_effect_download_is_unpacked(self):
        import zipfile
        with tempfile.TemporaryDirectory() as workspace:
            target = self.downloads(workspace)
            write_effect(target, [["1001", "甲", "生效中", "0.9"]])
            archive = target / "effect.zip"
            with zipfile.ZipFile(archive, "w") as bundle:
                bundle.write(target / "effect.csv", arcname="运营效果.csv")
            (target / "effect.csv").unlink()
            header, rows = W.read_effect(archive)
            self.assertEqual(len(rows), 1)
            self.assertIn("经营任务完成率", header)


class ZeroHardGate(unittest.TestCase):
    """0 行 / 0 组 硬闸：绝不生成、绝不发送。"""

    def test_zero_rows_is_no_report(self):
        result = W.deliver({"rows": 0, "path": None, "stats": []}, Path("/tmp"), "2026-10-13")
        self.assertEqual(result["status"], "no-report")
        self.assertEqual(result["sendAttempts"], 0)

    def test_zero_stats_is_no_report(self):
        result = W.deliver({"rows": 10, "path": "/x.xlsx", "stats": [],
                            "statsPath": "/x.png", "name": "a.xlsx", "statsName": "b.png"},
                           Path("/tmp"), "2026-10-13")
        self.assertEqual(result["status"], "no-report")
        self.assertEqual(result["sendAttempts"], 0)

    def test_disarmed_blocks_before_touching_gateway(self):
        original = W.DISARM_FILE
        with tempfile.TemporaryDirectory() as workspace:
            W.DISARM_FILE = Path(workspace) / "SEND_DISARMED.json"
            W.DISARM_FILE.write_text("{}", encoding="utf-8")
            try:
                result = W.deliver({"rows": 5, "path": "/x.xlsx", "stats": [{"name": "刘欢", "avg": 90.0}],
                                    "statsPath": "/x.png", "name": "a.xlsx", "statsName": "b.png"},
                                   Path(workspace), "2026-10-13")
            finally:
                W.DISARM_FILE = original
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["reason"], "send-disarmed")


class StatsCard(unittest.TestCase):
    def test_color_thresholds(self):
        self.assertEqual(stats_card.value_color(89.99), stats_card.COLOR_RED)
        self.assertEqual(stats_card.value_color(90.0), stats_card.COLOR_BLACK)
        self.assertEqual(stats_card.value_color(95.0), stats_card.COLOR_BLACK)
        self.assertEqual(stats_card.value_color(95.01), stats_card.COLOR_GREEN)

    def test_number_format_matches_sample(self):
        self.assertEqual(stats_card.format_pct(92.08), "92.08")
        self.assertEqual(stats_card.format_pct(91.8), "91.8")
        self.assertEqual(stats_card.format_pct(88.9), "88.9")
        self.assertEqual(stats_card.format_pct(88.0), "88")

    def test_display_name_cleaning(self):
        self.assertEqual(stats_card.display_name("盯个盯徐从众"), "徐从众")
        self.assertEqual(stats_card.display_name("上海青辅咨询"), "青辅")
        self.assertEqual(stats_card.display_name("张燕茹发数据"), "张燕茹")
        self.assertEqual(stats_card.display_name("未知新名字"), "未知新名字")

    def test_display_name_never_blank(self):
        for raw in ("", "   ", "咨询", "科技"):
            self.assertTrue(stats_card.display_name(raw).strip())

    def test_distinct_owners_never_merge_even_with_same_display_name(self):
        """★ 防回归（10-13 真实事故）：曾把「雷洁老师」和「雷洁2」都清洗成「雷洁」，
        按展示名分组时两家门店混算成一条 95.33——那 5 家全是雷洁2 的，
        而雷洁老师的 23 家贡献 0 家并静默消失。分组键必须永远是原始运营人。"""
        rows = [("雷洁老师", 70.0), ("雷洁2", 96.6667), ("雷洁2", 93.3333)]
        stats = stats_card.compute_stats(rows)
        by_raw = {item["raw"]: item for item in stats}
        self.assertEqual(set(by_raw), {"雷洁老师", "雷洁2"})
        self.assertAlmostEqual(by_raw["雷洁2"]["avg"], 95.0)
        self.assertAlmostEqual(by_raw["雷洁老师"]["avg"], 70.0)
        self.assertEqual(by_raw["雷洁2"]["shops"], 2)
        self.assertEqual(by_raw["雷洁老师"]["shops"], 1)

    def test_collision_falls_back_to_raw_name(self):
        """清洗后撞名的两个原始运营人，各自独立成行、用原始值显示，绝不合并。"""
        rows = [("王森", 90.0), ("王森组长号", 99.0)]
        stats = stats_card.compute_stats(rows)
        self.assertEqual(sorted(item["name"] for item in stats), ["王森", "王森组长号"])
        self.assertAlmostEqual(next(i["avg"] for i in stats if i["raw"] == "王森"), 90.0)
        self.assertAlmostEqual(next(i["avg"] for i in stats if i["raw"] == "王森组长号"), 99.0)

    def test_lei_teacher_display_name_stays_distinct_from_lei2(self):
        self.assertEqual(stats_card.display_name("雷洁老师"), "雷洁")
        self.assertEqual(stats_card.display_name("雷洁2"), "雷洁2")
        self.assertNotEqual(stats_card.display_name("雷洁老师"), stats_card.display_name("雷洁2"))

    def test_sort_is_descending(self):
        rows = [("甲", 80.0), ("乙", 95.0), ("丙", 88.0)]
        self.assertEqual([item["name"] for item in stats_card.compute_stats(rows)], ["乙", "丙", "甲"])

    def test_two_char_name_padded_to_three_cells(self):
        self.assertEqual(stats_card.pad_name("王森", 3.0), "王　森")
        self.assertEqual(stats_card.pad_name("青辅", 3.0), "青　辅")
        self.assertEqual(stats_card.pad_name("岳贝贝", 3.0), "岳贝贝")

    def test_render_writes_white_background_png(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as workspace:
            out = stats_card.render_card(stats_card.compute_stats([("王森", 92.0), ("青辅", 88.0)]),
                                         Path(workspace) / "card.png")
            self.assertTrue(out.exists())
            image = Image.open(out).convert("RGB")
            self.assertEqual(image.getpixel((2, 2)), stats_card.COLOR_BG)  # 白底
            self.assertGreater(image.width, 100)


class CoverageNotes(unittest.TestCase):
    """图尾提示行：「<原始运营人> N 家无数据未计入」（大同 10-13 口径）。"""

    def test_format_matches_sample(self):
        notes = stats_card.coverage_notes({"allocOnlyGroups": [
            {"raw": "雷洁老师", "shops": 23}]})
        self.assertEqual(notes, ["雷洁老师 23 家无数据未计入"])

    def test_empty_coverage_means_no_tail_line(self):
        self.assertEqual(stats_card.coverage_notes(None), [])
        self.assertEqual(stats_card.coverage_notes({}), [])
        self.assertEqual(stats_card.coverage_notes({"allocOnlyGroups": []}), [])

    def test_uses_raw_name_not_cleaned_display(self):
        """★ 必须用原始运营人：清洗会把「雷洁老师」压成「雷洁」，和图里的「雷洁2」难分。"""
        notes = stats_card.coverage_notes({"allocOnlyGroups": [
            {"raw": "雷洁老师", "shops": 23}, {"raw": "王森组长号", "shops": 1}]})
        self.assertEqual(notes, ["雷洁老师 23 家无数据未计入", "王森组长号 1 家无数据未计入"])

    def test_zero_shop_entry_is_skipped(self):
        notes = stats_card.coverage_notes({"allocOnlyGroups": [{"raw": "某组", "shops": 0}]})
        self.assertEqual(notes, [])

    def test_render_notes_append_visible_tail(self):
        """带提示行时：图变高、且底部区域有墨迹（无提示行时底部应是纯白）。"""
        from PIL import Image
        stats = stats_card.compute_stats([("王森", 92.0)])
        with tempfile.TemporaryDirectory() as workspace:
            plain = stats_card.render_card(stats, Path(workspace) / "plain.png")
            noted = stats_card.render_card(stats, Path(workspace) / "noted.png",
                                           notes=["雷洁老师 23 家无数据未计入"])
            plain_image = Image.open(plain).convert("RGB")
            noted_image = Image.open(noted).convert("RGB")
            self.assertGreater(noted_image.height, plain_image.height)
            # noted：底部 150px 必须圈到提示行墨迹（提示行位于倒数第二段）
            band = noted_image.crop((0, noted_image.height - 150, noted_image.width,
                                     noted_image.height))
            self.assertTrue(any(pixel != stats_card.COLOR_BG
                                for pixel in band.getdata()))
            # plain：底部 60px 无提示行，必须纯白
            plain_band = plain_image.crop((0, plain_image.height - 60, plain_image.width,
                                           plain_image.height))
            self.assertFalse(any(pixel != stats_card.COLOR_BG
                                 for pixel in plain_band.getdata()))


if __name__ == "__main__":
    unittest.main(verbosity=2)