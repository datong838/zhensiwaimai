#!/usr/bin/env python3
"""把 docs/metrics_catalog.json 同步进台账的『指标字典』sheet。

设计原则
--------
1. **单一权威源**：20 项指标只维护 `docs/metrics_catalog.json`，本脚本负责渲染进 Excel。
   改指标不要手改 xlsx，改 JSON 后重跑本脚本（可重复执行，幂等）。
2. **不重建工作簿**：只清空『指标字典』A2:H21 的值再写回，其它 sheet（店铺池/指标进度/
   当月反馈记录，含真实 486 行数据）原样保留。用 openpyxl 定点改值，不动样式与数据校验。
3. **零假数据**：JSON 里没有的 Mxx 行一律写『（待大同补发）』+ 状态『待补』。

用法
----
    python scripts/sync_metrics_catalog.py                        # 默认台账（当月）
    python scripts/sync_metrics_catalog.py --ledger docs/全店铺运营指标商家反馈台账_2026-10.xlsx
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re

from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent
WF = HERE.parent
CATALOG = WF / "docs" / "metrics_catalog.json"
TOTAL_METRICS = 20
PENDING = "（待大同补发）"
# 已拆出至其它工作流的指标：行内标注去向，其余列清空（防残留旧值误导）
# 来源=docs/metrics_catalog.json 的 split_out_metrics 字段（2026-10-22 01:13 拆流）


def default_ledger() -> pathlib.Path:
    """默认写『当月实盘台账』，绝不碰结构模板。

    坑（2026-10-20 踩过）：直接 glob + sorted()[-1] 会选中 `..._模板.xlsx`
    （中文『模』> 数字『2』，排序在后），把指标字典写进了模板。故显式排除『模板』，
    且只认 `_YYYY-MM.xlsx` 形式的实盘文件，多份时取修改时间最新的一份。
    """
    allx = list((WF / "docs").glob("全店铺运营指标商家反馈台账_*.xlsx"))
    cands = [p for p in allx if "模板" not in p.name and re.search(r"_\d{4}-\d{2}\.xlsx$", p.name)]
    if not cands:
        raise SystemExit(
            "未找到当月实盘台账（docs/全店铺运营指标商家反馈台账_YYYY-MM.xlsx）。"
            f"当前 docs 下候选：{[p.name for p in allx]}"
        )
    return max(cands, key=lambda p: p.stat().st_mtime)


def build_page_cell(m: dict) -> str:
    """把菜单路径/tab/路由拼成『所属页面／地址』一栏。"""
    parts = [m.get("menu_path") or ""]
    if m.get("tab"):
        parts.append(f"TAB：{m['tab']}")
    route = m.get("route")
    if route:
        parts.append(route)
    else:
        parts.append("路由待登录后取证")
    return " ｜ ".join(p for p in parts if p)


def build_notes_cell(m: dict) -> str:
    notes = m.get("capture_notes") or []
    marks = "①②③④⑤⑥⑦⑧⑨"
    out = []
    for i, n in enumerate(notes):
        mk = marks[i] if i < len(marks) else f"({i + 1})"
        out.append(f"{mk}{n}")
    return " ".join(out)


def sync_one(ledger: pathlib.Path, by_id: dict, split_out: dict | None = None) -> int:
    """把字典写进某一个工作簿，返回实填条数。表头不符拒绝（防写错文件）。"""
    wb = load_workbook(ledger)
    ws = wb["指标字典"]
    expect = ["序号", "指标名称", "所属页面／地址", "指标形态", "截图口径", "时间范围", "采集要点", "状态"]
    header = [c.value for c in ws[1]]
    if header[:8] != expect:
        raise SystemExit(f"表头与预期不符，拒绝写入。实际={header[:8]}　文件={ledger.name}")

    filled = 0
    split_n = 0
    for row in range(2, 2 + TOTAL_METRICS):
        mid = f"M{row - 1:02d}"
        m = by_id.get(mid)
        sp = split_out.get(mid)
        if sp:
            # 拆出行：标注去向 + 全列清空（openpyxl 写 None 不清空旧值，必须显式 ""）
            vals = [mid, f"（已拆出至 {sp.get('flow')}）", "", "", "", "", "", "已拆出"]
            split_n += 1
        elif not m:
            vals = [mid, PENDING, "", "", "", "", "", "待补"]
        else:
            vals = [
                mid,
                m.get("name"),
                build_page_cell(m),
                m.get("form"),
                m.get("shot"),
                m.get("time_range"),
                build_notes_cell(m),
                m.get("status", "已入字典"),
            ]
            filled += 1
        for col, v in enumerate(vals, start=1):
            ws.cell(row=row, column=col, value=v)

    ws.freeze_panes = "C2"  # 序号/名称两列锁住，横向滚动时还能对上店和指标
    wb.save(ledger)
    return filled


def template_path() -> pathlib.Path | None:
    tpl = [p for p in (WF / "docs").glob("全店铺运营指标商家反馈台账_*.xlsx") if "模板" in p.name]
    return tpl[0] if tpl else None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ledger", type=pathlib.Path, default=None)
    ap.add_argument("--catalog", type=pathlib.Path, default=CATALOG)
    ap.add_argument(
        "--no-template",
        action="store_true",
        help="只写实盘台账，不同步空白模板（默认两份都写，防止模板字典长期空白）",
    )
    args = ap.parse_args()

    ledger = args.ledger or default_ledger()
    cat = json.loads(args.catalog.read_text(encoding="utf-8"))
    by_id = {m["id"]: m for m in cat.get("metrics", [])}
    split_out = cat.get("split_out_metrics", {})

    wb = load_workbook(ledger)
    ws = wb["指标字典"]

    header = [c.value for c in ws[1]]
    expect = ["序号", "指标名称", "所属页面／地址", "指标形态", "截图口径", "时间范围", "采集要点", "状态"]
    if header[:8] != expect:
        raise SystemExit(f"表头与预期不符，拒绝写入。实际={header[:8]}")

    filled = 0
    split_n = 0
    for row in range(2, 2 + TOTAL_METRICS):
        mid = f"M{row - 1:02d}"
        m = by_id.get(mid)
        sp = split_out.get(mid)
        if sp:
            # 拆出行：标注去向 + 全列清空（openpyxl 写 None 不清空旧值，必须显式 ""）
            vals = [mid, f"（已拆出至 {sp.get('flow')}）", "", "", "", "", "", "已拆出"]
            split_n += 1
        elif not m:
            vals = [mid, PENDING, "", "", "", "", "", "待补"]
        else:
            vals = [
                mid,
                m.get("name"),
                build_page_cell(m),
                m.get("form"),
                m.get("shot"),
                m.get("time_range"),
                build_notes_cell(m),
                m.get("status", "已入字典"),
            ]
            filled += 1
        for col, v in enumerate(vals, start=1):
            ws.cell(row=row, column=col, value=v)

    ws.freeze_panes = "C2"  # 序号/名称两列锁住，横向滚动时还能对上店和指标
    wb.save(ledger)

    print(f"ledger  : {ledger}")
    print(f"catalog : {args.catalog}")
    print(f"已入字典: {filled}/{TOTAL_METRICS} 项；已拆出 {split_n} 项；"
          f"其余 {TOTAL_METRICS - filled - split_n} 项标『待补』")

    # 模板也要带字典：模板是给人看结构的，字典空着等于「模板没维护好」（2026-10-20 大同追问暴露）
    tpl = None if args.no_template else template_path()
    if tpl and tpl.resolve() != ledger.resolve():
        sync_one(tpl, by_id, split_out)
        print(f"template: {tpl}（已同步 {filled}/{TOTAL_METRICS}）")

    for mid in sorted(by_id):
        print(f"  {mid}  {by_id[mid].get('menu_path')} ｜ {by_id[mid].get('name')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
