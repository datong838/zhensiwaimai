"""Deterministic M01–M11 orchestration. Browser and sender are boundary adapters."""
from __future__ import annotations
import contextlib
import fcntl
import hashlib
import io
import json
import os
import re
import time
import uuid
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook
import ledger_ops as ledger_api

IDS = tuple(f'M{i:02}' for i in range(1, 12))
STOP = {'verified','already_verified','unknown','attempted','skipped-already'}
class WorkflowError(RuntimeError): pass

def atomic(path, value):
 path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
 tmp=path.with_name(path.name+'.'+uuid.uuid4().hex+'.tmp')
 with tmp.open('w',encoding='utf-8') as f:
  json.dump(value,f,ensure_ascii=False,indent=2);f.flush();os.fsync(f.fileno())
 os.replace(tmp,path)

@contextlib.contextmanager
def locked(path):
 path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
 with path.open('a+') as f:
  try:fcntl.flock(f,fcntl.LOCK_EX|fcntl.LOCK_NB)
  except BlockingIOError:raise WorkflowError('workflow-busy')
  try:yield
  finally:fcntl.flock(f,fcntl.LOCK_UN)

def select_metrics(day, config, explicit=None):
 mids=list(explicit or config.get('metric_ids',IDS))
 if not mids or len(set(mids))!=len(mids) or any(m not in IDS for m in mids):raise WorkflowError('metric-not-registered')
 if explicit:return mids
 start=config.get('start_date')
 if not start:raise WorkflowError('rotation-start-date-required')
 offset=(day-date.fromisoformat(start)).days
 if offset<0:raise WorkflowError('rotation-not-started')
 return [mids[offset%len(mids)]]

def validate_capture(r,sid,name,mid,run_dir):
 if not r.get('ok') or str(r.get('shopId'))!=sid or r.get('topName')!=name or r.get('metric')!=mid:raise WorkflowError('capture-identity-mismatch')
 if r.get('scopeVerified') is not True or r.get('timeVerified') is not True:raise WorkflowError('capture-scope-or-time-unverified')
 if r.get('truncated') or any(s.get('truncated') for s in r.get('shots',[])):raise WorkflowError('capture-truncated')
 files=r.get('files') or []
 if len(files)!=(3 if mid=='M11' else 1):raise WorkflowError('capture-image-count')
 from PIL import Image
 for raw in files:
  p=Path(raw).resolve()
  if not p.is_relative_to(Path(run_dir).resolve()) or not p.is_file() or p.stat().st_size<8000:raise WorkflowError('capture-file-invalid')
  with Image.open(p) as im:
   if im.format!='PNG' or min(im.size)<100:raise WorkflowError('capture-image-invalid')
   im.verify()
 return r

def compose(files,out):
 if len(files)==1:return Path(files[0])
 from PIL import Image
 ims=[]
 try:
  for f in files:
   with Image.open(f) as im:ims.append(im.convert('RGB'))
  canvas=Image.new('RGB',(max(i.width for i in ims),sum(i.height for i in ims)),'white');y=0
  for im in ims:canvas.paste(im,(0,y));y+=im.height
  out=Path(out);canvas.save(out,'PNG');canvas.close();return out
 finally:
  for im in ims:im.close()

class Workflow:
 def __init__(self,*,root,ledger,day,metrics,pool=None,group_map=None,capture=None,send=None,probe=None,throttle=2,max_age=7):
  self.root=Path(root);self.ledger=Path(ledger);self.day=day;self.month=day.strftime('%Y-%m')
  self.metrics=select_metrics(day,{},metrics);self.capture=capture;self.send=send;self.probe=probe;self.throttle=throttle
  self.ledger_keys={};self.pool=[]
  wb=load_workbook(self.ledger,data_only=False)
  try:
   self.ledger_keys=ledger_api.existing_keys(wb[ledger_api.RECORD_SHEET],self.month)
   ws=wb[ledger_api.POOL_SHEET];head=[str(c.value or '').strip() for c in ws[1]]
   def index(*names):
    for n in names:
     if n in head:return head.index(n)
    raise WorkflowError('pool-schema-invalid')
   ix=[index('店铺ID','店铺id'),index('店铺名称'),index('营业状态','店铺状态'),index('企业微信群名称'),index('采集日期')]
   for row in ws.iter_rows(min_row=2,values_only=True):
    if row[ix[0]] is None:continue
    self.pool.append(dict(zip(['id','name','status','group','collectedAt'],[row[i] for i in ix])))
  finally:wb.close()
  if pool is not None:self.pool=pool
  seen=set()
  for s in self.pool:
   s['id']=ledger_api.norm_id(s['id']);s['name']=str(s.get('name') or '').strip()
   if not re.fullmatch(r'\d+',s['id']) or not s['name'] or s['id'] in seen:raise WorkflowError('pool-identity-invalid')
   seen.add(s['id'])
   t=s.get('collectedAt');t=t.date() if isinstance(t,datetime) else t
   try:d=t if isinstance(t,date) else date.fromisoformat(str(t)[:10])
   except (ValueError,TypeError):raise WorkflowError('pool-date-invalid')
   if not 0<=(day-d).days<=max_age:raise WorkflowError('pool-snapshot-expired')
   s['group']=str((group_map.get(s['id'],'') if group_map is not None else s.get('group','')) or '').strip()
  self.pool=[s for s in self.pool if s.get('status')=='营业中']
  if not self.pool:raise WorkflowError('pool-empty')

 def event(self,sid,mid):return self.root/'events'/f'{self.month}-metrics-{sid}-{mid}.json'
 def checkpoint(self,sid,mid):return self.root/'checkpoints'/f'{self.month}-metrics-{sid}-{mid}.json'
 def prior(self,sid,mid):
  found=[]
  for folder in ['events','checkpoints']:
   for p in (self.root/folder).glob(f'{self.month}*-metrics-{sid}-{mid}.json'):
    try:found.append(json.loads(p.read_text()))
    except (ValueError,OSError):raise WorkflowError('event-unreadable')
  ledger=self.ledger_keys.get((sid,mid))
  if ledger in STOP:found.append({'state':'verified' if ledger in {'verified','skipped-already'} else ledger,'ledgerOnly':True})
  # Only a verified event bound to the exact attempted payload can reconcile a crash checkpoint.
  verified=[x for x in found if x.get('state')=='verified' and x.get('sha256')]
  found=[x for x in found if not (x.get('state')=='attempted' and any(v.get('sha256')==x.get('sha256') and v.get('group',v.get('groupName'))==x.get('group',x.get('groupName')) for v in verified))]
  # Uncertain dispatch always outranks an unbound success projection.
  for states in [('unknown','attempted'),('verified','already_verified','skipped-already')]:
   for r in found:
    if r.get('state') in states:return r
  return {}

 def record(self,shop,mid,prior,result='verified'):
  if prior.get('ledgerOnly'):return
  args=SimpleNamespace(ledger=self.ledger,metric=mid,date=self.day.isoformat(),store_id=shop['id'],store_name=shop['name'],group=prior.get('group',shop['group']),result=result,round='',receipt=prior.get('receipt',''),shot=prior.get('image',''),note='',require_group=True,force=False,apply=True)
  with locked(self.ledger.with_suffix('.xlsx.lock')),contextlib.redirect_stdout(io.StringIO()):
   if ledger_api.cmd_record(args)!=0:raise WorkflowError('ledger-write-failed')

 def run(self,mode='plan',authorization=''):
  if mode not in {'plan','capture','execute'}:raise WorkflowError('invalid-mode')
  if mode=='execute' and not authorization.strip():raise WorkflowError('authorization-ref-required')
  if mode=='plan':return self._run(mode,authorization)
  with locked(self.root/'batch.lock'):return self._run(mode,authorization)

 def _run(self,mode,authorization):
  started=time.monotonic();run_id=self.day.isoformat()+'-'+uuid.uuid4().hex[:12];run_dir=self.root/'runs'/run_id
  rows=[];halt=False
  for shop in self.pool:
   for mid in self.metrics:
    t=time.monotonic();row={'shopId':shop['id'],'shopName':shop['name'],'metric':mid,'group':shop['group'],'status':'pending','steps':[],'pitfalls':[]}
    def stage(label,fn):
     st=time.monotonic()
     try:return fn()
     finally:row['steps'].append({'stage':label,'elapsedMs':round((time.monotonic()-st)*1000)})
    prior=self.prior(shop['id'],mid);state=prior.get('state')
    if state in STOP:
     row['status']='unknown' if state in {'unknown','attempted'} else 'already_verified'
     if mode!='plan' and row['status']=='already_verified':
      try:stage('ledger-repair',lambda:self.record(shop,mid,prior))
      except Exception:row['status']='ledger_pending';row['pitfalls'].append('verified-but-ledger-write-failed')
    elif halt:row['status']='blocked';row['pitfalls'].append('batch-stopped-login-or-safety')
    elif not shop['group'] and mode!='capture':row['status']='blocked';row['pitfalls'].append('group-unresolved')
    elif mode!='plan':
     try:
      if not self.capture:raise WorkflowError('capture-adapter-missing')
      dest=run_dir/shop['id']/mid;dest.mkdir(parents=True,exist_ok=True)
      r=stage('capture',lambda:self.capture(shop,mid,dest))
      if r.get('status')=='no-data':
       row['status']='no-data'
       if self.probe is None or not stage('post-login-check',self.probe):raise WorkflowError('login-check-failed')
       if mode=='execute':stage('ledger-no-data',lambda:self.record(shop,mid,{},result='no-data'))
      else:
       validate_capture(r,shop['id'],shop['name'],mid,dest)
       image=stage('compose',lambda:compose(r['files'],dest/'combined.png'));row['image']=str(image)
       if self.probe is None or not stage('post-login-check',self.probe):raise WorkflowError('login-check-failed')
       if mode=='capture':row['status']='captured'
       else:
        if not self.send:raise WorkflowError('send-adapter-missing')
        info={'state':'attempted','group':shop['group'],'shopId':shop['id'],'metric':mid,'month':self.month,'image':str(image),'sha256':hashlib.sha256(image.read_bytes()).hexdigest(),'authorizationRef':authorization,'at':datetime.now().isoformat()}
        atomic(self.checkpoint(shop['id'],mid),info)
        try:sent=stage('delivery',lambda:self.send(shop,mid,image,self.day,authorization))
        except Exception:
         sent={'status':'unknown','reason':'sender-outcome-unknown'}
         ep=self.event(shop['id'],mid)
         if ep.exists():
          try:
           proof=json.loads(ep.read_text())
           if proof.get('state')=='verified' and proof.get('sha256')==info['sha256'] and proof.get('groupName',proof.get('group'))==shop['group']:sent={'status':'verified','receipt':proof.get('receipt','')}
          except (OSError,ValueError):pass
        status=sent.get('status','unknown')
        if status=='blocked' and sent.get('sendAttempts',0):status='unknown'
        state='verified' if status in {'verified','already_verified'} else 'blocked' if status=='blocked' else 'unknown'
        info.update(state=state,receipt=sent.get('receipt',''),result=sent);atomic(self.event(shop['id'],mid),info);atomic(self.checkpoint(shop['id'],mid),info)
        row['status']=status if state=='verified' else state
        if state=='verified':
         try:stage('ledger-write',lambda:self.record(shop,mid,info))
         except Exception:row['status']='ledger_pending';row['pitfalls'].append('verified-but-ledger-write-failed')
        if self.throttle:stage('throttle',lambda:time.sleep(self.throttle))
     except Exception as exc:
      reason=str(exc);row['status']='failed';row['pitfalls'].append(reason if re.fullmatch('[a-z0-9-]+',reason) else 'capture-or-validation-failed')
      if any(k in reason.lower() for k in ['login','captcha','security','risk']):halt=True
      elif self.probe:
       try:stage('failure-login-check',self.probe)
       except Exception:halt=True;row['pitfalls'].append('post-login-check-failed')
    row['elapsedMs']=round((time.monotonic()-t)*1000);rows.append(row)
    if mode!='plan':atomic(run_dir/'checkpoint.json',{'runId':run_id,'rows':rows})
  summary={'poolTotal':len(self.pool),'expected':len(rows)}
  for status in ['pending','blocked','unknown','already_verified','verified','captured','no-data','failed','ledger_pending']:summary[status]=sum(r['status']==status for r in rows)
  summary['remaining']=summary['expected']-summary['verified']-summary['already_verified']
  complete=summary['remaining']==0 if mode=='execute' else summary['captured']+summary['already_verified']==summary['expected'] if mode=='capture' else summary['blocked']+summary['unknown']==0
  result={'ok':complete,'workflowId':'daily-store-metrics-rotation-to-wechat','mode':mode,'status':'planned' if mode=='plan' else 'completed' if complete else 'incomplete','businessDate':self.day.isoformat(),'runId':run_id,'metrics':self.metrics,'summary':summary,'rows':rows,'elapsedMs':round((time.monotonic()-started)*1000),'externalEffectRequested':mode=='execute','stage':'metrics-completed' if complete else 'metrics-incomplete'}
  if mode!='plan':atomic(run_dir/'receipt.json',result)
  return result
