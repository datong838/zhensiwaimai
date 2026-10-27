import json
from datetime import date
from pathlib import Path
import pytest
from openpyxl import Workbook
from rotation_workflow import Workflow, select_metrics, validate_capture, WorkflowError

METRICS=[f'M{i:02}' for i in range(1,12)]

def ledger(path):
 from ledger_ops import RECORD_COLS
 w=Workbook();w.active.title='店铺池';w.active.append(['店铺ID','店铺名称','营业状态','运营组织','企业微信群名称','采集日期','备注'])
 w.active.append(['101','甲店','营业中','组织','甲群','2026-10-21',''])
 w.active.append(['102','乙店','营业中','组织','乙群','2026-10-21',''])
 s=w.create_sheet('当月反馈记录');s.append(RECORD_COLS)
 s=w.create_sheet('指标进度');s.append(['店铺id','店铺名称','企业微信群名称']+METRICS+['已完成数']);s.append(['101','甲店','甲群']);s.append(['102','乙店','乙群']);w.save(path)

@pytest.fixture
def setup(tmp_path):
 p=tmp_path/'ledger.xlsx';ledger(p)
 return tmp_path,p

def make(setup, **kw):
 root,p=setup
 return Workflow(root=root/'state',ledger=p,day=date(2026,10,21),metrics=kw.pop('metrics',['M01']),**kw)

def test_rotation_is_continuous_across_month_and_only_registered():
 assert select_metrics(date(2026,11,1),{'start_date':'2026-10-21','metric_ids':METRICS})==['M01']
 with pytest.raises(WorkflowError):select_metrics(date(2026,10,21),{'start_date':None,'metric_ids':METRICS})
 with pytest.raises(WorkflowError):select_metrics(date(2026,10,21),{'metric_ids':['M12'],'start_date':'2026-10-21'})

def test_plan_is_offline_and_covers_all_eleven(setup):
 def fail(*a):raise AssertionError('plan must not contact browser or sender')
 w=make(setup,metrics=METRICS,capture=fail,send=fail)
 p=w.run('plan');assert p['summary']['expected']==22;assert p['summary']['pending']==22
 assert not (w.root/'events').exists()

def test_attempted_unknown_and_cross_day_verified_never_capture_or_resend(setup):
 w=make(setup,capture=lambda *a:pytest.fail('must not capture'),send=lambda *a:pytest.fail('must not send'))
 w.root.mkdir();(w.root/'events').mkdir()
 for sid,state in [('101','attempted'),('102','verified')]:
  (w.root/'events'/f'2026-10-20-metrics-{sid}-M01.json').write_text(json.dumps({'state':state}))
 r=w.run('execute',authorization='test');assert not r['ok'];assert r['summary']['unknown']==1;assert r['summary']['already_verified']==1

def test_missing_group_is_blocked_before_capture(setup):
 w=make(setup,group_map={'101':'甲群'},capture=lambda *a:pytest.fail('plan-only'),send=lambda *a:None)
 p=w.run('plan');assert p['summary']['blocked']==1;assert p['summary']['pending']==1

def test_capture_rejects_wrong_scope_name_truncation_and_missing_time(tmp_path):
 from PIL import Image
 p=tmp_path/'shot.png';Image.effect_noise((700,400),100).convert('RGB').save(p)
 good={'ok':True,'metric':'M02','shopId':'101','topName':'甲店','scopeVerified':True,'timeVerified':True,'files':[str(p)],'truncated':False}
 assert validate_capture(good,'101','甲店','M02',tmp_path)['files']==[str(p)]
 for patch in [{'shopId':'102'},{'topName':'乙店'},{'truncated':True},{'timeVerified':False},{'scopeVerified':False}]:
  with pytest.raises(WorkflowError):validate_capture(good|patch,'101','甲店','M02',tmp_path)

def test_partial_batch_resumes_without_resending_verified(setup):
 calls=[]
 def cap(shop,mid,out):
  from PIL import Image
  out.mkdir(parents=True,exist_ok=True);p=out/'shot.png';Image.effect_noise((700,400),100).convert('RGB').save(p)
  return {'ok':True,'metric':mid,'shopId':shop['id'],'topName':shop['name'],'scopeVerified':True,'timeVerified':True,'files':[str(p)],'truncated':False}
 def send(shop,mid,image,day,auth):
  calls.append(shop['id']);return {'status':'verified','receipt':'test-only','sendAttempts':1}
 w=make(setup,capture=cap,send=send,probe=lambda:True,throttle=0)
 assert w.run('execute',authorization='test')['ok']
 assert len(calls)==2
 assert w.run('execute',authorization='test')['ok'];assert len(calls)==2

def test_sender_exception_is_unknown_not_retryable(setup):
 def cap(shop,mid,out):
  from PIL import Image
  out.mkdir(parents=True,exist_ok=True);p=out/'shot.png';Image.effect_noise((700,400),100).convert('RGB').save(p)
  return {'ok':True,'metric':mid,'shopId':shop['id'],'topName':shop['name'],'scopeVerified':True,'timeVerified':True,'files':[str(p)]}
 def boom(*a):raise TimeoutError('send timed out')
 w=make(setup,capture=cap,send=boom,probe=lambda:True,throttle=0)
 r=w.run('execute',authorization='test');assert not r['ok'];assert r['summary']['unknown']==2
 w.capture=lambda *a:pytest.fail('unknown cannot recapture');r=w.run('execute',authorization='test');assert r['summary']['unknown']==2

def test_all_eleven_flow_capture_send_and_replay_zero_send(setup):
 calls=[]
 def cap(shop,mid,out):
  from PIL import Image
  files=[]
  for i in range(3 if mid=='M11' else 1):
   p=out/f'{i}.png';Image.effect_noise((700,400),100).convert('RGB').save(p);files.append(str(p))
  return {'ok':True,'metric':mid,'shopId':shop['id'],'topName':shop['name'],'scopeVerified':True,'timeVerified':True,'files':files}
 def send(shop,mid,image,*a):
  from PIL import Image
  with Image.open(image) as im:assert im.height==(1200 if mid=='M11' else 400)
  calls.append((shop['id'],mid));return {'status':'verified','receipt':'isolated-test','sendAttempts':1}
 w=make(setup,metrics=METRICS,capture=cap,send=send,probe=lambda:True,throttle=0)
 r=w.run('execute',authorization='isolated-test');assert r['ok'];assert len(calls)==22
 assert all(row['steps'] and row['elapsedMs']>=0 for row in r['rows'])
 assert w.run('execute',authorization='isolated-test')['ok'];assert len(calls)==22

def test_verified_but_ledger_failed_repairs_without_resending(setup,monkeypatch):
 from PIL import Image
 sent=[]
 def cap(shop,mid,out):
  p=out/'a.png';Image.effect_noise((700,400),100).convert('RGB').save(p)
  return {'ok':True,'metric':mid,'shopId':shop['id'],'topName':shop['name'],'scopeVerified':True,'timeVerified':True,'files':[str(p)]}
 w=make(setup,capture=cap,send=lambda *a:(sent.append(1) or {'status':'verified'}),probe=lambda:True,throttle=0)
 original=w.record
 monkeypatch.setattr(w,'record',lambda *a:(_ for _ in ()).throw(OSError('disk full')))
 assert w.run('execute',authorization='test')['summary']['ledger_pending']==2
 monkeypatch.setattr(w,'record',original);w.capture=lambda *a:pytest.fail('must only repair ledger')
 assert w.run('execute',authorization='test')['ok'];assert len(sent)==2

def test_login_failure_stops_remaining_shops_and_empty_data_never_sends(setup):
 w=make(setup,capture=lambda *a:(_ for _ in ()).throw(WorkflowError('brand-login-required')),send=lambda *a:pytest.fail('no send'),probe=lambda:True,throttle=0)
 r=w.run('capture');assert r['summary']['failed']==1;assert r['summary']['blocked']==1
 w.capture=lambda *a:{'status':'no-data'};r=w.run('capture');assert r['summary']['no-data']==2

def test_expired_pool_and_duplicate_shop_fail_closed(setup):
 pool=[{'id':'1','name':'A','status':'营业中','group':'G','collectedAt':'2026-01-01'}]
 with pytest.raises(WorkflowError,match='expired'):make(setup,pool=pool)
 pool[0]['collectedAt']='2026-10-21'
 with pytest.raises(WorkflowError,match='identity'):make(setup,pool=pool*2)

def test_verified_event_can_reconcile_exact_crash_checkpoint_only(setup):
 w=make(setup);w.root.mkdir();(w.root/'events').mkdir();(w.root/'checkpoints').mkdir()
 for key,sha in [('events','same'),('checkpoints','same')]:
  state='verified' if key=='events' else 'attempted'
  (w.root/key/'2026-10-metrics-101-M01.json').write_text(json.dumps({'state':state,'sha256':sha,'group':'甲群'}))
 assert w.prior('101','M01')['state']=='verified'
 (w.root/'checkpoints'/'2026-10-metrics-101-M01.json').write_text(json.dumps({'state':'attempted','sha256':'other','group':'甲群'}))
 assert w.prior('101','M01')['state']=='attempted'

def test_official_entry_plan_is_offline_and_limit_does_not_claim_full_pool(setup,capsys,monkeypatch):
 from metrics_rotation_daily import main
 import subprocess
 root,p=setup
 monkeypatch.setattr(subprocess,'run',lambda *a,**k:pytest.fail('offline plan must not spawn browser or sender'))
 assert main(['--ledger',str(p),'--runtime-dir',str(root/'state'),'--date','2026-10-21','--metrics','M01,M11','--limit','1'])==0
 r=json.loads(capsys.readouterr().out);assert r['sourcePoolTotal']==2;assert r['scopeLimited'];assert r['summary']['expected']==2

def test_unknown_ledger_never_marks_completed_progress(setup,capsys):
 from types import SimpleNamespace
 from openpyxl import load_workbook
 import ledger_ops
 root,p=setup
 args=SimpleNamespace(ledger=p,metric='M01',date='2026-10-21',store_id='101',store_name='甲店',group='甲群',result='unknown',round='',receipt='',shot='',note='',require_group=True,force=False,apply=True)
 assert ledger_ops.cmd_record(args)==0
 w=load_workbook(p);assert w['指标进度'].cell(2,4).value is None;assert w['当月反馈记录'].max_row==2;w.close()

def test_capture_does_not_require_group_and_cannot_send_or_mark_feedback(setup):
 from PIL import Image
 def cap(shop,mid,out):
  p=out/'shot.png';Image.effect_noise((700,400),100).convert('RGB').save(p)
  return {'ok':True,'metric':mid,'shopId':shop['id'],'topName':shop['name'],'scopeVerified':True,'timeVerified':True,'files':[str(p)]}
 w=make(setup,group_map={},capture=cap,send=lambda *a:pytest.fail('capture must never call sender'),probe=lambda:True)
 before=w.ledger.read_bytes();r=w.run('capture')
 assert r['ok'];assert r['summary']['captured']==2;assert w.ledger.read_bytes()==before
 assert not (w.root/'events').exists()
