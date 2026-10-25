import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import run_all
class LedgerDeliveryTests(unittest.TestCase):
    def test_business_fact_skips_without_rewriting_success(self):
        from datetime import datetime
        with patch('business_delivery_facts.read_fact',return_value={'businessDelivered':True}),patch.object(run_all,'_persist_done_today_record') as persist:
            self.assertEqual(run_all._cadence_gate({'name':'confirmed-flow'},datetime.now()),
                             (True,'business_delivery_observed'))
            persist.assert_not_called()
    def test_failed_delivery_never_marks_sent(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(run_all,'send_ledger_via_wechat',return_value=False):
            path=Path(folder)/'sent.json'
            with patch.object(run_all,'SENT_TODAY',path):run_all._maybe_send_ledger(Path(folder)/'report.xlsx',True,True,'2026-10-24')
            self.assertFalse(path.exists())
    def test_wrong_frozen_path_never_sends(self):
        with patch('authorized_supplement.load_grant',return_value={'item':{'path':'/fixed/report.xlsx'}}),patch('authorized_supplement.deliver_once') as send:
            self.assertFalse(run_all.send_ledger_via_wechat(Path('/other/report.xlsx')));send.assert_not_called()
    def test_verified_original_delivery_marks_sent(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(run_all,'send_ledger_via_wechat',return_value=True):
            path=Path(folder)/'sent.json'
            with patch.object(run_all,'SENT_TODAY',path):run_all._maybe_send_ledger(Path(folder)/'report.xlsx',True,True,'2026-10-24')
            self.assertTrue(path.exists())
class CustomerLedgerViewTests(unittest.TestCase):
    def rows(self):
        return [{'name': str(i), 'exec_time': '', 'note': '旧技术异常或历史说明', 'done': done,
                 'dev_state': 'paused' if done=='已合并' else 'active', 'superseded':done=='不补跑'}
                for i,done in enumerate(['是','已合并','不补跑'])]
    def test_tomorrow_uses_yes_and_only_current_day_notes(self):
        rows=self.rows();before=[dict(r) for r in rows]
        shown,total,done=run_all.customer_ledger_view(rows,'2026-10-26')
        self.assertEqual([r['done'] for r in shown],['是']*3)
        self.assertEqual([r['note'] for r in shown],['当日已完成','已随当日申诉完成','由当日末班覆盖'])
        self.assertEqual((total,done),(3,3));self.assertEqual(rows,before)
    def test_today_retains_original_presentation(self):
        rows=self.rows();shown,total,done=run_all.customer_ledger_view(rows,'2026-10-25')
        self.assertEqual(shown,rows);self.assertEqual((total,done),(1,1))
    def test_unfinished_stays_unfinished_in_gate(self):
        rows=self.rows()+[{'name':'pending','done':'否','dev_state':'active','note':'未完成','exec_time':''}]
        shown,total,done=run_all.customer_ledger_view(rows,'2026-10-26')
        self.assertEqual(len(shown),3);self.assertEqual((total,done),(4,3))
    def test_excel_and_csv_customer_columns_match(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            csv=Path(folder)/'test.csv';xlsx=Path(folder)/'test.xlsx';rows=self.rows()
            run_all.write_ledger_csv(rows,csv,'2026-10-26')
            run_all.write_ledger_xlsx(rows,xlsx,'2026-10-26')
            text=csv.read_text();self.assertNotIn('旧技术异常',text)
            self.assertIn('应执行数量,3',text);self.assertIn('已完成数量,3',text)
            w=load_workbook(xlsx,read_only=True,data_only=True)
            actual=list(w.active.iter_rows(min_row=4,max_row=6,values_only=True));w.close()
            self.assertTrue(all(r[2]=='是' for r in actual))
            self.assertEqual([r[3] for r in actual],['当日已完成','已随当日申诉完成','由当日末班覆盖'])


class LedgerScheduleDisplayTests(unittest.TestCase):
    def test_current_manifest_retains_exact_original_twenty_one(self):
        import json
        data=json.loads((Path(__file__).parent/'skills_manifest.json').read_text())
        names=[r['skill'] for r in data['ledger_template']]
        self.assertEqual(len(names),21)
        self.assertEqual(len(set(names)),21)
        self.assertEqual(names[0],'daily-agent-new-merchants-to-wechat')
        self.assertEqual(names[-1],'meituan-intelligent-customer-service')
        self.assertIn('taobao-shangou-negative-review-appeals',names)
        self.assertIn('meituan-waimai-negative-review-appeals-1425',names)
    def test_csv_retains_merge_but_not_uncompleted_business(self):
        rows=[{'name':str(i),'schedule':'','exec_time':'','note':'','done':done,
               'dev_state':'paused' if done=='已合并' else 'active'}
              for i,done in enumerate(['是','已合并','不补跑','否'])]
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ledger.csv'
            run_all.write_ledger_csv(rows,path,'2026-10-25')
            text=path.read_text()
            self.assertIn('1,,已合并,',text)
            self.assertIn('2,,不补跑,',text)
            self.assertNotIn('3,,否,',text)
    def rows(self, skill, record=None):
        data={'skills':[skill], 'ledger_template':[{'row':1,'skill':skill['name'],
            'ledger_name':'业务任务','schedule':'旧时间','dev_state':'active'}]}
        with patch('business_delivery_facts.read_fact',return_value=None):
            return run_all.build_ledger_rows(data, {skill['name']:record} if record else {})
    def test_current_schedule_is_authoritative(self):
        row=self.rows({'name':'daily-negative-reviews','enabled':True,'schedule':'每日08:30'})[0]
        self.assertEqual(row['schedule'],'每日08:30')
        self.assertEqual(row['done'],'否')
    def test_explicit_merged_slot_is_not_active_or_success(self):
        row=self.rows({'name':'taobao-shangou-negative-review-appeals','enabled':False,
                       'schedule':'已合并到下午四档，不单独排班'})[0]
        self.assertEqual(row['done'],'已合并')
        self.assertEqual(row['dev_state'],'paused')
        self.assertEqual(row['exec_time'],'')
    def test_disabled_without_merge_is_not_completed(self):
        row=self.rows({'name':'unready','enabled':False,'schedule':'等待恢复'})[0]
        self.assertEqual(row['done'],'否')
        self.assertEqual(row['dev_state'],'active')
    def test_real_failure_retains_incomplete_status(self):
        row=self.rows({'name':'daily-taobao-download-to-jinsha','enabled':True,'schedule':'每日13:00'},
                      {'category':'UNKNOWN','status':'FAILED'})[0]
        self.assertEqual(row['done'],'否')
        self.assertEqual(row['dev_state'],'active')

if __name__=='__main__':unittest.main()
