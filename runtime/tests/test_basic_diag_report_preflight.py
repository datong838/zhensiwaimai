import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
FILE=ROOT/'skills/workflows/daily-store-basic-diagnosis-to-wechat/scripts/basic_diag_workflow.py'
spec=importlib.util.spec_from_file_location('basic_diag_preflight_original',FILE)
flow=importlib.util.module_from_spec(spec)
spec.loader.exec_module(flow)

class BasicReportPreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.run=Path(self.temp.name);self.downloads=self.run/'downloads';self.downloads.mkdir()
        self.events=self.run/'events';self.events.mkdir()
        p=patch.object(flow,'EVENTS',self.events);p.start();self.addCleanup(p.stop)
        p=patch.object(flow,'arm_state',return_value={'armed':True});p.start();self.addCleanup(p.stop)
        self.day='2026-10-28'
        self.alloc([['门店ID','运营人'],[' 123 ','测试运营']])
    def alloc(self,rows):
        wb=openpyxl.Workbook();ws=wb.active
        for row in rows:ws.append(row)
        wb.save(self.downloads/'alloc.xlsx');wb.close()
    def diagnosis(self,metrics):
        (self.downloads/'diagnosis.json').write_text(json.dumps({'metrics':metrics},ensure_ascii=False))
    def sample(self,status='未完成',sid=123):
        return {'shop_id':sid,'shop_name':'测试门店','task_name':'测试问题','task_status':status,'current_value':'1','target_value':'2'}
    def report(self):
        self.diagnosis({'指标A':{'rows':[self.sample()]},'指标B':{'rows':[self.sample(sid=999)]},'零行指标':{'rows':[]}})
        return flow.build_report(self.run,self.day)
    def test_original_join_and_one_sheet_per_nonempty_metric(self):
        result=self.report();self.assertEqual((result['rows'],result['matched'],result['unmatched']),(2,1,1))
        wb=openpyxl.load_workbook(result['path']);self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames,['指标A','指标B'])
        self.assertEqual([c.value for c in wb['指标A'][1]],flow.COLUMNS)
        self.assertEqual(wb['指标A'].cell(2,6).value,'测试运营')
        self.assertEqual(wb['指标B'].cell(2,6).value,'未分配')
    def test_zero_all_metrics_no_file(self):
        self.diagnosis({'指标A':{'rows':[]}})
        result=flow.build_report(self.run,self.day)
        self.assertEqual(result['rows'],0);self.assertIsNone(result['path'])
    def test_completed_tasks_not_in_report(self):
        self.diagnosis({'指标A':{'rows':[self.sample('已完成')]}})
        self.assertEqual(flow.build_report(self.run,self.day)['rows'],0)
    def test_missing_alloc_blocks(self):
        self.diagnosis({});(self.downloads/'alloc.xlsx').unlink()
        with self.assertRaisesRegex(ValueError,'alloc-download-missing'):flow.build_report(self.run,self.day)
    def test_missing_alloc_columns_blocks(self):
        self.diagnosis({});self.alloc([['门店ID','其他列'],['123','运营']])
        with self.assertRaisesRegex(ValueError,'alloc-missing-required-columns'):flow.build_report(self.run,self.day)
    def test_verified_marker_no_adapter_or_resend(self):
        report=self.report();marker=self.events/f'{self.day}-store-basic-diagnosis-ops.json'
        marker.write_text(json.dumps({'state':'verified','sha256':'prior-payload-sha','groupName':flow.GROUP}))
        original=marker.read_bytes()
        with patch.dict('sys.modules',{'wecom_adapter':None}):result=flow.deliver(report,self.run,self.day)
        self.assertEqual(result['status'],'already_verified');self.assertEqual(result['sendAttempts'],0)
        self.assertEqual(marker.read_bytes(),original)
    def test_unknown_marker_readonly(self):
        report=self.report();marker=self.events/f'{self.day}-store-basic-diagnosis-ops.json'
        marker.write_text(json.dumps({'state':'unknown','sha256':hashlib.sha256(Path(report['path']).read_bytes()).hexdigest()}))
        original=marker.read_bytes()
        with patch.dict('sys.modules',{'wecom_adapter':None}):result=flow.deliver(report,self.run,self.day)
        self.assertEqual(result['status'],'unknown');self.assertEqual(result['sendAttempts'],0)
        self.assertEqual(marker.read_bytes(),original)
    def test_no_report_does_not_touch_adapter(self):
        with patch.dict('sys.modules',{'wecom_adapter':None}):result=flow.deliver({'rows':0},self.run,self.day)
        self.assertEqual(result,{'status':'no-report','sendAttempts':0})

if __name__=='__main__':unittest.main()
