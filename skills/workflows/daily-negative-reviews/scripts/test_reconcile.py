import json,tempfile,unittest
from pathlib import Path
import openpyxl
import reconcile

class ReconcileTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.date='2026-09-17'
        self.raw=self.root/'raw.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='差评分析数据';sheet.append([]);sheet.append([]);sheet.append(reconcile.FIELDS);sheet.append(['门店A',2,self.date,'淘宝闪购','',None,'订单A']);book.save(self.raw)
        self.ui=self.root/'ui.json';self.ui.write_text(json.dumps({'counts':{'淘宝闪购':1},'rows':[{'store':'门店A','score':'总分2','date':self.date,'platform':'淘宝闪购','comment':'','status':'商家已隐藏评价'}],'businessWarnings':['授权失败']},ensure_ascii=False))
        self.config={'raw':str(self.raw),'uiStatus':str(self.ui),'businessDate':self.date,'finalData':str(self.root/'final.json'),'analysis':str(self.root/'analysis.json')}
    def tearDown(self):self.temp.cleanup()
    def test_reconcile_and_supplement(self):
        result=reconcile.reconcile(self.config);self.assertTrue(result['ok']);self.assertEqual(result['profile']['records'],1);self.assertEqual(result['profile']['supplementedFromUi'],1);self.assertEqual(json.loads(Path(self.config['finalData']).read_text())[0]['差评内容'],'商家已隐藏评价')
    def test_mismatch_fails_closed(self):
        value=json.loads(self.ui.read_text());value['rows'][0]['comment']='不一致';self.ui.write_text(json.dumps(value,ensure_ascii=False))
        with self.assertRaisesRegex(ValueError,'multiset'):reconcile.reconcile(self.config)
    def test_zero_reviews_produces_valid_empty_dataset(self):
        book=openpyxl.Workbook();sheet=book.active;sheet.title='差评分析数据';sheet.append([]);sheet.append([]);sheet.append(reconcile.FIELDS);book.save(self.raw)
        self.ui.write_text(json.dumps({'counts':{'美团外卖':0,'淘宝闪购':0,'京东外卖':0},'rows':[],'businessWarnings':[]},ensure_ascii=False))
        result=reconcile.reconcile(self.config);self.assertEqual(result['profile']['records'],0);self.assertEqual(json.loads(Path(self.config['finalData']).read_text()),[])
    def test_above_three_is_audited_and_excluded(self):
        book=openpyxl.load_workbook(self.raw);sheet=book.active;sheet.append(['门店B',4,self.date,'美团外卖','口味可以的',None,None]);book.save(self.raw)
        value=json.loads(self.ui.read_text());value['counts']['美团外卖']=1;value['rows'].append({'store':'门店B','score':'总分4','date':self.date,'platform':'美团外卖','comment':'口味可以的','status':''});self.ui.write_text(json.dumps(value,ensure_ascii=False))
        result=reconcile.reconcile(self.config);rows=json.loads(Path(self.config['finalData']).read_text())
        self.assertEqual(result['profile']['sourceRecords'],2);self.assertEqual(result['profile']['records'],1);self.assertEqual(result['profile']['scoresAbove3'],0);self.assertEqual(result['profile']['excludedAboveThree'],1);self.assertEqual([row['店铺名称'] for row in rows],['门店A'])

if __name__=='__main__':unittest.main()
